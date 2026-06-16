import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Notifications Framework"

headers = [
    "Group",
    "Source",
    "Trigger Logic",
    "Update vs. Alert",
    "Where User Is Notified",
    "Where User Can Resolve",
    "Where Data Is Stored",
    "For How Long",
    "Clarification Flag / Notes",
]

rows = [
    ["Interactions Report", "Interactions Report",
     "Interaction Completed (user is attended); Interaction Report created to log notes",
     "Alert", "Activity Feed, Teams, Int. Hub, Int. Report",
     "Data Source",
     "Activity Feed\nTeams Homepage\nInteractions Report\nInteractions Hub Calendar view",
     "90 days past date of alert/update\n1 month \"past due\" after date of interaction\nIndefinitely\n1 week past date of notes logged",
     ""],

    ["Interactions Report", "Interactions Report",
     "Interaction Completed (user not attended; no deal team); Interaction Report created to log notes",
     "Update", "Activity Feed, Teams, Int. Hub, Int. Report",
     "Data Source",
     "Activity Feed\nTeams Homepage\nInteractions Report\nInteractions Hub Calendar view",
     "90 days past date of alert/update\n1 month past due after date of interaction\nIndefinitely\n1 week past date of notes logged",
     ""],

    ["Interactions Report", "Interactions Report",
     "Upcoming Tasks in next 3 days (user has been assigned the task)",
     "Alert", "Activity Feed",
     "Data Source",
     "Activity Feed\nTeams Homepage\nInteractions Report\nInteractions Hub Tasks tab",
     "90 days past date of alert/update\n1 month past due after date of interaction\nIndefinitely\n1 month past due date",
     ""],

    ["Interactions Report", "Interactions Report",
     "Upcoming Tasks in next 3 days (user assigned the task to someone else)",
     "Update", "Activity Feed",
     "Data Source",
     "Activity Feed\nTeams Homepage\nInteractions Report\nInteractions Hub Tasks tab",
     "90 days past date of alert/update\n1 month past due after date of interaction\nIndefinitely\n1 month past due date",
     "[C4-Row4] Retention inferred from \"Upcoming Task\" family (Row 3) — not given separately in feedback. Confirm or provide distinct retention."],

    ["Interactions Report", "Interactions Report",
     "A task you assigned to another user has been deleted",
     "Update", "Activity Feed",
     "Data Source",
     "Activity Feed\nTeams Homepage\nInteractions Report (shows as deleted)\nInteractions Hub Tasks tab",
     "90 days past date of alert/update\n1 month past due after date of interaction\nIndefinitely\n1 week past date of update",
     ""],

    ["Opportunity Details / Opportunities", "Opportunity Details / Opportunities",
     "Banker updates a key deal attribute, such as (but not limited to) probability, coverage, estimated close date, or estimated fee.",
     "Update", "Activity Feed",
     "Data Source",
     "Activity Feed\nOpportunity Details\nOpportunities widget\nChange Log\nWeekly Change Digest",
     "90 days past date of alert/update\n1 week past date of alert/update\n1 week past date of alert/update\n1 year past date of alert/update\n1 month past date of alert/update",
     ""],

    ["Opportunity Details / Opportunities", "Opportunity Details / Opportunities",
     "Change in Deal State to Closed-Won",
     "Alert", "Activity Feed",
     "Data Source",
     "Activity Feed\nTextbot\nOpportunity Details\nOpportunities widget\nChange Log\nWeekly Change Digest",
     "90 days past date of alert/update\n1 month past date of alert/update\n1 week past date of alert/update\n1 week past date of alert/update\n1 year past date of alert/update\n1 month past date of alert/update",
     ""],

    ["Opportunity Details / Opportunities", "Opportunity Details",
     "When new opportunity is created or coverage users updated to include banker they are alerted",
     "Update", "Activity Feed",
     "Data Source",
     "Activity Feed\nOpportunity Details\nOpportunities widget\nChange Log\nWeekly Change Digest",
     "90 days past date of alert/update\n1 week past date of alert/update\n1 week past date of alert/update\n1 year past date of alert/update\n1 month past date of alert/update",
     ""],

    ["Engagement Health", "Engagement Health",
     "Banking's last interaction date is coming due in, or My Last Interaction Date is coming due — both within 2 weeks",
     "Alert", "Activity Feed",
     "Data Source",
     "Activity Feed\nEngagement Health\nInteractions page",
     "90 days past date of alert/update\n1 year \"past due\" (after which client marked 'inactive')\nIndefinitely (after 1 year past due, client becomes 'inactive')",
     ""],

    ["Signals", "Signals",
     "News item flagged for action",
     "Alert", "Activity Feed",
     "Data Source",
     "Activity Feed\nSignals",
     "90 days past date of alert/update\n90 days",
     ""],

    ["Approvals Queue", "Approvals Queue",
     "When another banker approves a client tier change, for example, that this banker's client tier change has been approved",
     "Update", "Operations",
     "Data Source",
     "Activity Feed\nOpportunity Details",
     "90 days past date of alert/update\n1 week past date of update",
     "[C4-Row11] Trigger logic copy not explicitly updated in feedback (only retention given) — left as original."],

    ["Approvals Queue", "Approvals Queue",
     "Another banker on same deal team requests a deal attribute change that requires approval.",
     "Update", "Operations",
     "Data Source",
     "Activity Feed",
     "90 days past date of alert/update",
     "Type changed from Alert to Update per feedback."],

    ["Hierarchy Cases", "Hierarchy Cases",
     "When another banker approves a client legal entity change, for example, that this banker requested for an account they cover",
     "Update", "Operations",
     "Data Source",
     "Activity Feed\nOpportunity Details",
     "90 days past date of alert/update\n1 week past date of update",
     ""],

    ["Hierarchy Cases", "Hierarchy Cases",
     "Another banker on same deal team requests a change to the status of a covered account which requires approval.",
     "Update", "Operations",
     "Data Source",
     "Activity Feed",
     "90 days past date of alert/update",
     "Type changed from Alert to Update per feedback."],
]

header_fill = PatternFill(start_color="1A1A1A", end_color="1A1A1A", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=10)
for col_idx, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(wrap_text=True, vertical="center")

note_font = Font(color="B45309", italic=True, size=9)
wrap_align = Alignment(wrap_text=True, vertical="top")

row_idx = 2
for row in rows:
    for col_idx, value in enumerate(row, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = wrap_align
        if col_idx == 9 and value:
            cell.font = note_font
    row_idx += 1

widths = [22, 26, 50, 14, 28, 14, 30, 38, 50]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 30
for r in range(2, row_idx):
    ws.row_dimensions[r].height = 90

notes_ws = wb.create_sheet("Notes & Legend")
notes_ws["A1"] = "How to use this sheet"
notes_ws["A1"].font = Font(bold=True, size=12)
notes_ws["A3"] = "1. Edit any copy/text cells directly in the 'Notifications Framework' tab."
notes_ws["A4"] = "2. Do not rename column headers — they map directly to the HTML structure."
notes_ws["A5"] = "3. 'Where Data Is Stored' and 'For How Long' are paired line-by-line (line 1 of each column corresponds to line 1 of the other, etc). Keep them in sync if you add/remove a line."
notes_ws["A6"] = "4. 'Where User Can Resolve' is now a single value (Data Source) for every row per latest feedback."
notes_ws["A7"] = "5. The 'Clarification Flag / Notes' column lists open questions still requiring confirmation."
notes_ws["A8"] = "6. Send this file back and changes will be applied to notifications-framework.html."
notes_ws.column_dimensions["A"].width = 100
for r in range(3, 9):
    notes_ws[f"A{r}"].alignment = Alignment(wrap_text=True)

wb.save("/home/user/claude-code-workspace/notifications-framework.xlsx")
print("saved")
